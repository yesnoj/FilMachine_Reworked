# -*- coding: utf-8 -*-
"""Manufacturing outputs for JLCPCB (PCB + SMT/THT assembly): Gerbers + drill (zip), BOM CSV, CPL CSV with the
JLCPCB rotation conventions, schematic/assembly PDFs, STEP, renders.  Usage: export.py <out_dir>"""
import sys, os, re, subprocess, csv, zipfile, collections, shutil, json, math
import pcbnew
import design

OUT = sys.argv[1]
PCB = f"{OUT}/FilMachine_Mainboard.kicad_pcb"
SCH = f"{OUT}/FilMachine_Mainboard.kicad_sch"
FAB = f"{OUT}/fab"
ROT_DB = "/home/claude/jlc/cpl_rotations_db.csv"
shutil.rmtree(FAB, ignore_errors=True)
os.makedirs(f"{FAB}/gerber", exist_ok=True)

def run(*a):
    r = subprocess.run(list(a), capture_output=True, text=True)
    if r.returncode:
        print("ERR", " ".join(a[:4]), r.stdout[-500:], r.stderr[-500:])
    return r

# ── Gerbers (Protel extensions, no X2 attributes, silk minus mask) + Excellon drill (PTH+NPTH in one file) ──
LAYERS = "F.Cu,In1.Cu,In2.Cu,B.Cu,F.Paste,B.Paste,F.Silkscreen,B.Silkscreen,F.Mask,B.Mask,Edge.Cuts"
run("kicad-cli", "pcb", "export", "gerbers", "-o", f"{FAB}/gerber/", "-l", LAYERS,
    "--no-x2", "--no-netlist", "--subtract-soldermask", PCB)
run("kicad-cli", "pcb", "export", "drill", "-o", f"{FAB}/gerber/", "--format", "excellon", "--drill-origin", "absolute",
    "--excellon-units", "mm", "--excellon-zeros-format", "decimal", "--generate-map", "--map-format", "gerberx2", PCB)
z = zipfile.ZipFile(f"{FAB}/FilMachine_Mainboard_Gerber.zip", "w", zipfile.ZIP_DEFLATED)
for fn in sorted(os.listdir(f"{FAB}/gerber")):
    z.write(f"{FAB}/gerber/{fn}", fn)
z.close()

# ── CPL: Designator, Mid X, Mid Y, Layer, Rotation ─────────────────────────────────────────────────
# JLCPCB places each part with ITS OWN library footprint (the EasyEDA footprint of the LCSC part): "Mid X/Y" is that
# footprint's origin and "Rotation" is measured from that footprint's 0 deg orientation (counter-clockwise).
# jlcpos.py compares every EasyEDA footprint with the KiCad footprint used here and stores, per KiCad footprint and
# LCSC code, the rotation offset and the position of the EasyEDA origin in KiCad footprint coordinates
# (jlcpos.json). Without an entry we fall back to the KiCad anchor + the public rotation table.
_JP = [p for p in (os.path.join(os.path.dirname(os.path.abspath(__file__)), "jlcpos.json"), "/home/claude/jlc/jlcpos.json") if os.path.exists(p)]
JLCPOS = json.load(open(_JP[0])) if _JP else {}
rot_rules = []
for row in csv.reader(open(ROT_DB)):
    if not row or row[0].startswith("Footprint pattern") or row[0].startswith("#"):
        continue
    rot_rules.append((re.compile(row[0]), float(row[1])))
def db_rotation(fpname):
    for rx, r in rot_rules:
        if rx.search(fpname): return r
    return 0.0
def R(x, y, th):
    c, s_ = math.cos(math.radians(th)), math.sin(math.radians(th))
    return (x * c + y * s_, -x * s_ + y * c)       # KiCad: positive = CCW on screen (y down)

board = pcbnew.LoadBoard(PCB)
cpl_rows = []
fallback = collections.OrderedDict()
for fp in board.GetFootprints():
    if fp.IsDNP() or fp.IsExcludedFromPosFiles():
        continue
    if fp.GetLayer() != pcbnew.F_Cu:
        raise SystemExit("bottom-side part: " + fp.GetReference())
    fpid = str(fp.GetFPID().GetLibNickname()) + ":" + str(fp.GetFPID().GetLibItemName())
    lcsc = fp.GetFieldByName("LCSC").GetText() if fp.GetFieldByName("LCSC") else ""
    krot = fp.GetOrientationDegrees()
    pos = fp.GetPosition(); ax, ay = pos.x / 1e6, pos.y / 1e6
    ent = JLCPOS.get(f"{fpid}|{lcsc}")
    if ent and "error" not in ent:
        dx, dy = R(ent["tx"], ent["ty"], krot)
        cx, cy = ax + dx, ay + dy
        rot = (krot + ent["dtheta"]) % 360.0
    else:
        name = fpid.split(":")[1]
        cx, cy = ax, ay
        if fp.GetAttributes() & pcbnew.FP_THROUGH_HOLE and not fpid.startswith("FM:"):
            c = fp.GetCourtyard(pcbnew.F_CrtYd).BBox()
            cx, cy = (c.GetLeft() + c.GetRight()) / 2e6, (c.GetTop() + c.GetBottom()) / 2e6
        rot = (krot + db_rotation(name)) % 360.0
        fallback.setdefault(f"{fpid}|{lcsc}", []).append(fp.GetReference())
    # KiCad pos files negate Y (Gerber coordinates grow upwards); do the same so CPL and Gerbers agree
    cpl_rows.append((fp.GetReference(), round(cx, 4), round(-cy, 4), "Top", round(rot, 1)))
def natkey(r):
    m = re.match(r"([A-Z]+)(\d+)", r)
    return (m.group(1), int(m.group(2))) if m else (r, 0)
cpl_rows.sort(key=lambda r: natkey(r[0]))
with open(f"{FAB}/FilMachine_Mainboard_CPL.csv", "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["Designator", "Mid X", "Mid Y", "Layer", "Rotation"])
    for r in cpl_rows:
        w.writerow(r)
cpl_refs = {r[0] for r in cpl_rows}

# ── BOM: Comment, Designator, Footprint, LCSC Part # (one line per LCSC part; DNP and mechanical excluded) ──
groups = collections.OrderedDict()
for p in design.PARTS:
    if p["sym"].startswith(("Mechanical", "Jumper")) or p["dnp"]:
        continue
    if not p.get("lcsc"):
        raise SystemExit("no LCSC code: " + p["ref"])
    groups.setdefault(p["lcsc"], []).append(p)
with open(f"{FAB}/FilMachine_Mainboard_BOM.csv", "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["Comment", "Designator", "Footprint", "LCSC Part #"])
    for lcsc, ps in sorted(groups.items(), key=lambda kv: natkey(kv[1][0]["ref"])):
        vals = sorted({q["value"] for q in ps})
        fps = sorted({q["fp"].split(":")[1] for q in ps})
        assert len(fps) == 1, (lcsc, fps)
        w.writerow([" / ".join(vals), ",".join(sorted((q["ref"] for q in ps), key=natkey)), fps[0], lcsc])
bom_refs = {p["ref"] for ps in groups.values() for p in ps}
print("BOM lines:", len(groups), "placed refs:", len(bom_refs), "CPL refs:", len(cpl_refs),
      "missing in CPL:", sorted(bom_refs - cpl_refs), "extra in CPL:", sorted(cpl_refs - bom_refs))
print("CPL entries from the EasyEDA/JLCPCB footprint comparison:", len(cpl_rows) - sum(len(v) for v in fallback.values()),
      "| fallback (KiCad anchor + public rotation table):", {k: v for k, v in fallback.items()})

# ── a human-readable BOM too (xlsx): value, designators, MPN, manufacturer, description, LCSC ──
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BOM"
    ws.append(["Qty", "Designator", "Value", "Manufacturer", "MPN", "LCSC", "Footprint", "Description"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
    for lcsc, ps in sorted(groups.items(), key=lambda kv: natkey(kv[1][0]["ref"])):
        ws.append([len(ps), ",".join(sorted((q["ref"] for q in ps), key=natkey)), " / ".join(sorted({q["value"] for q in ps})),
                   ps[0]["mfr"], ps[0]["mpn"], lcsc, ps[0]["fp"].split(":")[1], ps[0]["desc"]])
    for i, w_ in enumerate([5, 40, 16, 20, 26, 12, 40, 70]):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w_
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # The DNP parts go on a separate sheet: a note row inside the BOM table would be parsed as a part if someone
    # uploaded this xlsx to JLCPCB ("R6,R7,J6 designators don't exist in the CPL file").
    dnp = [p for p in design.PARTS if p["dnp"]]
    if dnp:
        wn = wb.create_sheet("Non montati")
        wn.append(["Designator", "Value", "Note"])
        for c in wn[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
        for p in sorted(dnp, key=lambda q: natkey(q["ref"])):
            wn.append([p["ref"], p["value"], "DNP: non montato, non e' nella BOM/CPL JLCPCB. " + p.get("desc", "")])
        for i, w_ in enumerate([12, 16, 90]):
            wn.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w_
    wb.save(f"{FAB}/FilMachine_Mainboard_BOM_leggibile.xlsx")
except ImportError:
    pass

# ── documentation ──
run("kicad-cli", "sch", "export", "pdf", "-o", f"{FAB}/FilMachine_Mainboard_Schema.pdf", SCH)
run("kicad-cli", "pcb", "export", "pdf", "-o", f"{FAB}/FilMachine_Mainboard_Assembly_Top.pdf",
    "-l", "F.Fab,F.Silkscreen,Edge.Cuts", "--include-border-title", PCB)
run("kicad-cli", "pcb", "export", "step", "-o", f"{FAB}/FilMachine_Mainboard.step", "--subst-models", "--force", PCB)
# renders: 2D top/bottom (copper + silk) and a 3D view
def svgpng(layers, out, w=2400):
    svg = f"{FAB}/_r.svg"
    run("kicad-cli", "pcb", "export", "svg", "--mode-single", "--fit-page-to-board", "--exclude-drawing-sheet", "-l", layers, "-o", svg, PCB)
    run("rsvg-convert", "-w", str(w), "-b", "white", svg, "-o", out)
    if os.path.exists(svg): os.remove(svg)
svgpng("F.Cu,F.SilkS,F.Mask,Edge.Cuts", f"{FAB}/FilMachine_Mainboard_top.png")
svgpng("B.Cu,B.SilkS,Edge.Cuts", f"{FAB}/FilMachine_Mainboard_bottom.png")
run("xvfb-run", "-a", "kicad-cli", "pcb", "render", "-o", f"{FAB}/FilMachine_Mainboard_3D.png", "--side", "top",
    "--background", "opaque", "--quality", "high", "-w", "1800", "-h", "1500", "--zoom", "1.1", PCB)
print("export done:", sorted(os.listdir(FAB)))
