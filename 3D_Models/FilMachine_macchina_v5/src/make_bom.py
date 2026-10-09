"""Nuova distinta (docs/NuovaDistinta.xlsx): cose da acquistare, componenti gia' acquistati e loro uso, pezzi stampati, viteria."""
import os, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from common import OUT

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F"); GRP = PatternFill("solid", fgColor="E8EEF5"); TOT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 "€"'

def header(ws, row, cols, widths):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10); cell.fill = HDR; cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30

def put(ws, row, col, value, bold=False, color="000000", fmt=None, wrap=True, fill=None, align=None, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=F, bold=bold, color=color, size=size); c.border = BOX
    c.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal=align)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c

BUY = [  # gruppo, componente, specifiche, quantita', unita', prezzo unitario stimato, dove si usa, note
 ("Bagno e vaschette", "Contenitore Euro 400 x 300 x 170 mm", "polipropilene, pareti e fondo CHIUSI (non forato), maniglie chiuse; interno circa 365 x 265 x 165", 1, "pz", 15.0, "bagno termostatico", "regge 70 °C; vanno fatti due fori Ø16,5 per i riscaldatori"),
 ("Bagno e vaschette", "Vaschetta Gastronorm GN 1/9, altezza 150 mm", "176 x 108 mm, 1,5 l, polipropilene, con coperchio", 3, "pz", 8.0, "chimiche C1, C2, C3", "il coperchio serve per conservare la chimica fuori dalla macchina"),
 ("Bagno e vaschette", "Tubo quadro di alluminio 15 x 15 x 1,5 mm", "barra da 1 m (servono due spezzoni da 356 mm)", 1, "pz", 5.0, "longheroni che reggono le vaschette", ""),
 ("Bagno e vaschette", "Dado basso M16 inox + 2 guarnizioni piane in silicone 16 x 24 x 2", "una serie per riscaldatore", 2, "serie", 3.0, "passaggio dei riscaldatori nella parete del bagno", "controllare il passo del filetto sul riscaldatore"),
 ("Bagno e vaschette", "Termostato bimetallico KSD301 normalmente chiuso, 60 °C, 10 A", "", 2, "pz", 1.5, "sicurezza, in serie a ogni riscaldatore", "i riscaldatori non devono mai lavorare a secco"),
 ("Bagno e vaschette", "Nastro biadesivo per esterni (tipo VHB) 19 mm", "rotolo corto", 1, "pz", 3.0, "sensori di livello XKC-Y21 sulla parete del bagno", ""),
 ("Idraulica", "Tubo LLDPE 3/8\" (9,52 x 6,35 mm) per osmosi inversa", "rotolo da 10 m (ne servono circa 3)", 1, "pz", 8.0, "tutte le linee", ""),
 ("Idraulica", "Gomito a innesto rapido 3/8\" - 3/8\"", "per tubo da 9,52 mm", 10, "pz", 0.7, "6 nelle tre testine, 1 sopra la valvola dello scarico, 1 ai piedi della salita verso la tank", "2 di scorta; meglio in polipropilene"),
 ("Idraulica", "Gomito a codolo 3/8\" (un lato a innesto, l'altro a codolo liscio da 9,52 mm)", "in inglese \"stem elbow\"", 2, "pz", 1.5, "uscita del collettore (il codolo entra nella sede con O-ring) e raccordo destro della pompa", "codolo lungo almeno 33 mm dall'angolo sul collettore; sulla pompa va accorciato a 27 mm"),
 ("Idraulica", "Raccordo diritto: filetto maschio 3/8\" BSP - tubo 3/8\" a innesto", "", 2, "pz", 1.5, "attacchi della pompa", "con nastro in PTFE"),
 ("Idraulica", "Passaparete a innesto rapido 3/8\" - 3/8\"", "foro Ø16", 1, "pz", 2.0, "attacco SCARICO sulla staffa C09", ""),
 ("Idraulica", "Raccordo a T 3/8\" + rubinetto a sfera 3/8\" a innesto", "", 1, "serie", 4.0, "scarico manuale del bagno", "facoltativo"),
 ("Idraulica", "Tubo in silicone 8 x 12 mm, NERO", "3 m", 1, "pz", 9.0, "tre sifoni, tubo della tank, troppopieno, tubi verso le taniche", "nero: i due tubi che arrivano alla tank non devono far passare la luce"),
 ("Idraulica", "O-ring 9 x 2 in EPDM", "", 10, "pz", 0.2, "sei sedi del collettore C07", "4 di scorta"),
 ("Idraulica", "Fascette in nylon 2,5 x 100 mm", "confezione da 100", 1, "conf", 3.0, "valvole, sifoni, testine, cavi", ""),
 ("Idraulica", "Grasso al silicone per O-ring", "tubetto", 1, "pz", 5.0, "O-ring, paraolio, baionetta della spirale", ""),
 ("Tank e caricatore", "Albero in acciaio inox Ø8 x 61 mm", "barra rettificata h6/h7, due spianature per i grani", 1, "pz", 4.0, "trasmissione della spirale", ""),
 ("Tank e caricatore", "Cuscinetto 688-2RS (8 x 16 x 5)", "meglio in acciaio inox", 2, "pz", 2.5, "cartuccia B17", ""),
 ("Tank e caricatore", "Paraolio 8 x 16 x 5", "NBR", 1, "pz", 2.0, "tenuta dell'albero verso la vasca", ""),
 ("Tank e caricatore", "Rondella inox M8 (8,4 x 16 x 1,6)", "", 1, "pz", 0.2, "reggispinta tra trascinatore e cartuccia", ""),
 ("Tank e caricatore", "O-ring 24 x 2", "", 1, "pz", 0.3, "tenuta della flangia della cartuccia", ""),
 ("Tank e caricatore", "O-ring 80 x 3 in EPDM", "", 2, "pz", 1.0, "cinghia tonda tra motore e spirale", "uno di scorta"),
 ("Tank e caricatore", "Magneti al neodimio Ø6 x 3", "confezione da 10 (ne servono 4)", 1, "conf", 3.0, "ruota magneti del sensore Hall", ""),
 ("Tank e caricatore", "Molla a compressione Ø6 x 22 mm, filo 0,5", "", 1, "pz", 1.0, "perno folle a molla", ""),
 ("Tank e caricatore", "Micro-servo MG90S", "con squadretta a un braccio", 1, "pz", 4.0, "taglierina", ""),
 ("Tank e caricatore", "Lame di ricambio per cutter da 9 mm", "confezione da 10", 1, "conf", 2.0, "due spezzoni da 36 mm nella taglierina", ""),
 ("Tank e caricatore", "Velluto nero adesivo", "striscia da 1 m", 1, "pz", 4.0, "guarnizione di luce su coperchio e fessura", "facoltativo, da decidere dopo la prova di tenuta alla luce"),
 ("Elettrico", "Presa IEC C14 da pannello con interruttore e portafusibile", "foro 27,5 x 47,5 mm", 1, "pz", 4.0, "fianco sinistro del vano elettronica", ""),
 ("Elettrico", "Cavo di alimentazione con spina e connettore IEC C13", "", 1, "pz", 4.0, "", ""),
 ("Elettrico", "Modulo MOSFET 15 A (tipo D4184), ingresso 3,3-20 V", "", 2, "pz", 2.5, "riscaldatori (uscite B6 e B7)", "al posto del modulo IRF540, che a 8,3 A scalda troppo"),
 ("Elettrico", "Portafusibile a lama in linea con fusibile", "uno da 10 A, uno da 5 A", 2, "pz", 1.5, "pompa; valvole e logica", "consigliati"),
 ("Elettrico", "Cavo unipolare 1,5 mm2 rosso e nero", "5 m + 5 m", 1, "serie", 8.0, "linee a 12 V di riscaldatori e pompa", ""),
 ("Elettrico", "Cavetti per segnali, puntalini, connettori", "assortimento", 1, "serie", 5.0, "sensori, valvole, servo", ""),
 ("Elettrico", "Passacavo in gomma Ø22", "", 4, "pz", 0.5, "paratia, fianco destro, cavalletto", ""),
 ("Elettrico", "Resistenza 4,7 kohm", "", 1, "pz", 0.1, "pull-up del bus DS18B20", "se non e' gia' sulla scheda"),
 ("Elettrico", "Tester per servo", "", 1, "pz", 3.0, "prova della taglierina senza firmware", "facoltativo"),
 ("Struttura e viteria", "Multistrato 420 x 505 x 12 mm", "tagliato a misura", 1, "pz", 8.0, "base della macchina", "verniciare o impregnare"),
 ("Struttura e viteria", "Piedini in gomma Ø20", "", 4, "pz", 0.5, "sotto la base", ""),
 ("Struttura e viteria", "Viti autofilettanti per plastica 3 x 8", "testa cilindrica, confezione da 50 (ne servono 28)", 1, "conf", 4.0, "vedi foglio Viteria", ""),
 ("Struttura e viteria", "Viti autofilettanti per plastica 3 x 10", "testa cilindrica, confezione da 100 (ne servono 30)", 1, "conf", 5.0, "vedi foglio Viteria", ""),
 ("Struttura e viteria", "Viti autofilettanti per plastica 3 x 16", "testa cilindrica, confezione da 50 (ne servono 8)", 1, "conf", 4.0, "vedi foglio Viteria", ""),
 ("Struttura e viteria", "Viti M3 x 16 con dado", "20 pezzi (ne servono 10)", 1, "conf", 3.0, "morsetti delle schede", "M3 x 12 e' troppo corta: morsetto 8,7 mm + ripiano 3 mm + dado"),
 ("Struttura e viteria", "Minuteria: 3 grani M3 x 6 con dado M3, 4 viti autofilettanti M2 x 6, 1 vite M2,5 x 8, 2 viti M3 x 8, 2 viti M3 x 10, 1 vite M3 x 5", "assortimento", 1, "serie", 4.0, "trascinatore, pulegge, morsetti lama, manovella, motoriduttore, perno folle, modulo Hall", "una delle M3 x 10 fa da spinotto tra biella e portalama, al posto della spina Ø3 x 8"),
 ("Struttura e viteria", "Viti per lamiera 2,9 x 9,5 (8 pezzi) e 2,9 x 13 (6 pezzi)", "testa cilindrica; una confezione piccola per misura", 1, "serie", 3.0, "staffe, fermi e portasonde sui longheroni di alluminio", "forare l'alluminio a 2,4 mm"),
 ("Struttura e viteria", "Viti M4 x 12", "4 pezzi", 1, "conf", 1.0, "testate del telaio sotto la fascia del contenitore", ""),
 ("Struttura e viteria", "Viti da legno 3,5 x 12 a testa cilindrica (+ 4 da 3,5 x 20)", "confezione da 50", 1, "conf", 3.0, "tutti i pezzi fissati alla base (dima); le 4 lunghe per le colonnine", "non piu' lunghe: la base e' spessa 12 mm"),
 ("Struttura e viteria", "Barra filettata M3 x 130 mm con dadi e rondelle", "4 pezzi", 1, "serie", 3.0, "colonna dei tre ripiani a griglia", ""),
 ("Struttura e viteria", "Filamento PETG, bobina da 1 kg", "", 2, "pz", 20.0, "tutti i pezzi stampati", "stima: vedi foglio Pezzi stampati"),
]

OWNED = [  # componente, specifiche, quantita', prezzo, uso in questa macchina
 ("Alimentatore 12V 30A", "", 1, 28.18, "Usato: in piedi nello zoccolo D16, contro la paratia"),
 ("Display", "JC4880P443C", 1, 25.59, "Usato: sulla plancia inclinata D09"),
 ("Pompa", "Marco UP3-R", 1, 184.46, "Usata: sotto il cavalletto della tank. Tenerla a PWM basso e sciacquarla con acqua a fine lavoro (corpo in ottone nichelato)"),
 ("Elettrovalvola", "DC 12V Normally Close", 5, 7.59, "Usate tutte, in fila sul collettore C07. Da sinistra: C1, WB, C2, C3, scarico (i collegamenti elettrici non cambiano)"),
 ("Mosfet driver", "IRF540 V4.0", 1, 5.9, "NON usato per i riscaldatori: a 8,3 A per canale scalda troppo. Tenerlo di scorta"),
 ("Motor driver", "DC5-12V 0A-30A Dual channel H bridge", 1, 17.29, "Usato: canale A motore della spirale, canale B pompa"),
 ("Solenoid Driver", "MCP23017 GPIO", 1, 13.0, "Usato: valvole su A0-A4, riscaldatori su B6-B7"),
 ("Speaker", "Speaker 8ohm 2W", 1, 1.69, "Usato: dietro la griglia del fianco sinistro"),
 ("Motore agitatore", "JGB37-520 12V 20RPM", 1, 8.34, "Usato: nella galleria sotto il ripiano della tank, con cinghia tonda"),
 ("Portafusibile 15A", "Fusibile per riscaldatore", 2, 5.5, "Usati: uno per riscaldatore"),
 ("Riscaldatore", "100W 12V", 2, 10.85, "Usati: nella parete destra del bagno, vicino al fondo"),
 ("Convertitore 12V / 5V", "LM2596S", 1, 4.29, "Usato: scheda display, sensori di livello, servo"),
 ("Cavo Flat", "Cavo flat grigio 2x13 20cm", 1, 6.75, "Usato: dal display alla breakout"),
 ("Breakout board", "DC3 26 p 2x13Pins", 1, 10.34, "Usata: sul secondo ripiano"),
 ("Sensore hall", "KY-003 A3144", 1, 4.83, "Usato: sulla staffa B25, davanti alla ruota magneti"),
 ("Sensore flusso", "YF-S201 12V", 1, 5.94, "Non usato: serve solo con l'ingresso acqua in pressione"),
 ("Sensore livello", "XKC-Y21 5V", 8, 5.53, "Usati 2 (MIN e MAX del bagno); gli altri 6 restano di scorta: sulle vaschette estraibili non servono"),
 ("Sensore temperatura", "DS18B20 30mm 1m", 2, 3.31, "Usati: uno nel bagno (portasonde C06), uno nella vaschetta C1"),
]

SCREWS = [  # tipo, quantita', dove. Le lunghezze delle autofilettanti da 3 mm sono quelle verificate da check_screws.py
 ("Vite autofilettante 3 x 10", 4, "blocchi ponte sul ripiano della tank"), ("Vite autofilettante 3 x 16", 4, "tetto del ponte"),
 ("Vite autofilettante 3 x 10", 4, "culla universale"), ("Vite M2 x 6 autofilettante", 4, "morsetti delle lame"),
 ("Vite M2 x 8 (del servo)", 2, "alette del servo"), ("Vite M2,5 x 8", 1, "perno di manovella"), ("Spina Ø3 x 8 (o vite M3 x 10)", 1, "spinotto biella-portalama"),
 ("Vite autofilettante 3 x 8", 2, "coperchietto del vano servo"), ("Vite autofilettante 3 x 8", 2, "flangia della cartuccia"),
 ("Vite autofilettante 3 x 10", 1, "flangia della cartuccia, nel punto che regge anche la staffa Hall"),
 ("Vite autofilettante 3 x 10", 3, "piastra del motore"),
 ("Vite M3 x 8", 2, "motoriduttore sulla piastra: nel riduttore entrano 4 mm; se i suoi fori sono meno profondi usare M3 x 6"),
 ("Grano M3 x 6 + dado M3", 3, "trascinatore e due pulegge"), ("Vite M3 x 10", 1, "ritegno del perno folle"),
 ("Vite M3 x 5 (o una fascetta)", 1, "modulo Hall sulla staffa B25: piu' lunga toccherebbe il carter"),
 ("Vite autofilettante 3 x 16", 4, "piedini della tank sul cavalletto"),
 ("Vite per lamiera 2,9 x 9,5", 6, "staffe con ganci C03 sul longherone anteriore"), ("Vite autofilettante 3 x 10", 6, "culle del sifone C02 sulle staffe C03, da dietro"),
 ("Vite per lamiera 2,9 x 13", 6, "fermi posteriori C04 (spessi 7 mm)"), ("Vite per lamiera 2,9 x 9,5", 2, "portasonde C06"), ("Vite M4 x 12", 4, "testate del telaio C05"),
 ("Vite da legno 3,5 x 12", 2, "collettore C07"), ("Vite da legno 3,5 x 12", 6, "pettine valvole C08"), ("Vite da legno 3,5 x 12", 2, "staffa scarichi C09"),
 ("Vite da legno 3,5 x 12", 2, "carter riscaldatori C10"), ("Vite da legno 3,5 x 12", 6, "fianchi del cavalletto"),
 ("Vite da legno 3,5 x 12", 4, "piede della pompa: scegliere la lunghezza secondo i piedini, senza passare la base da 12 mm"),
 ("Vite autofilettante 3 x 10", 6, "piano del cavalletto sui fianchi"), ("Vite autofilettante 3 x 10", 4, "frontale del cavalletto"), ("Vite autofilettante 3 x 10", 2, "carter della trasmissione"),
 ("Vite da legno 3,5 x 12", 6, "fianchi del vano elettronica"), ("Vite autofilettante 3 x 8", 6, "frontale del vano elettronica"), ("Vite autofilettante 3 x 8", 6, "paratia posteriore"),
 ("Vite autofilettante 3 x 8", 4, "tetto"), ("Vite autofilettante 3 x 8", 4, "plancia del display"), ("Vite autofilettante 3 x 8", 4, "linguette D17 che tengono la scheda display"),
 ("Vite da legno 3,5 x 20", 4, "colonnine basse (ripiano inferiore)"), ("Vite da legno 3,5 x 12", 4, "zoccolo dell'alimentatore"),
 ("Barra filettata M3 x 130 + 2 dadi", 4, "colonna dei ripiani"), ("Vite M3 x 16 + dado", 10, "morsetti delle schede"),
]

def build():
    wb = Workbook()
    # ------------------------------------------------------------ Da acquistare
    ws = wb.active; ws.title = "Da acquistare"
    ws["A1"] = "FilMachine - nuova distinta: componenti da acquistare"; ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ("Prezzi unitari in blu: stime indicative (ottobre 2026), da verificare all'acquisto. Controllati sul web solo: contenitore Euro RS PRO 400 x 300 x 170 = 26,13 € IVA inclusa "
                "(it.rs-online.com, codice 163-1899; altrove costa meno) e vaschette Araven GN 1/9 150 mm con coperchio, circa 7,50 € l'una (nisbets.com.au, codice T983, confezione da 4).")
    ws["A2"].font = Font(name=F, size=9, italic=True); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells("A2:J2"); ws.row_dimensions[2].height = 40
    cols = ["N.", "Gruppo", "Componente", "Specifiche", "Q.tà", "Unità", "Prezzo unitario indicativo (€)", "Totale (€)", "Dove si usa", "Note"]
    header(ws, 4, cols, [5, 18, 40, 38, 7, 7, 13, 12, 38, 40])
    r0 = 5
    for i, (g, c, s, q, u, p, w, n) in enumerate(BUY):
        r = r0 + i
        put(ws, r, 1, i + 1, align="center"); put(ws, r, 2, g); put(ws, r, 3, c, bold=True); put(ws, r, 4, s)
        put(ws, r, 5, q, color="0000FF", align="center"); put(ws, r, 6, u, align="center"); put(ws, r, 7, p, color="0000FF", fmt=EUR)
        put(ws, r, 8, f"=E{r}*G{r}", fmt=EUR); put(ws, r, 9, w); put(ws, r, 10, n)
    r1 = r0 + len(BUY) - 1; rt = r1 + 1
    put(ws, rt, 3, "TOTALE da acquistare", bold=True, fill=TOT); put(ws, rt, 8, f"=SUM(H{r0}:H{r1})", bold=True, fmt=EUR, fill=TOT)
    for j in (1, 2, 4, 5, 6, 7, 9, 10): put(ws, rt, j, None, fill=TOT)
    rs = rt + 2
    put(ws, rs, 3, "Totale per gruppo", bold=True, fill=GRP); put(ws, rs, 8, None, fill=GRP)
    groups = []
    for g, *_ in BUY:
        if g not in groups: groups.append(g)
    for k, g in enumerate(groups, 1):
        put(ws, rs + k, 3, g); put(ws, rs + k, 8, f'=SUMIF($B${r0}:$B${r1},C{rs + k},$H${r0}:$H${r1})', fmt=EUR)
    rf = rs + len(groups) + 1
    put(ws, rf, 3, "di cui facoltativi (T e rubinetto, velluto, tester per servo)", bold=False)
    opt = [r0 + i for i, b in enumerate(BUY) if "facoltativo" in b[7]]
    put(ws, rf, 8, "=" + "+".join(f"H{r}" for r in opt), fmt=EUR)
    ws.freeze_panes = "D5"
    # ------------------------------------------------------------ Gia' acquistati
    ws2 = wb.create_sheet("Già acquistati")
    ws2["A1"] = "Componenti già acquistati (DistintaComponenti.xlsx) e loro uso in questa macchina"; ws2["A1"].font = Font(name=F, bold=True, size=14)
    ws2["A2"] = "Quantità e prezzi sono quelli della distinta originale del progetto."; ws2["A2"].font = Font(name=F, size=9, italic=True)
    header(ws2, 4, ["Componente", "Specifiche", "Q.tà", "Prezzo unitario (€)", "Totale (€)", "Uso in questa macchina"], [26, 36, 7, 13, 12, 90])
    for i, (c, s, q, p, u) in enumerate(OWNED):
        r = 5 + i
        put(ws2, r, 1, c, bold=True); put(ws2, r, 2, s); put(ws2, r, 3, q, align="center"); put(ws2, r, 4, p, fmt=EUR); put(ws2, r, 5, f"=C{r}*D{r}", fmt=EUR)
        put(ws2, r, 6, u, color="B3261E" if u.startswith(("NON", "Non")) else "000000")
    rt2 = 5 + len(OWNED)
    put(ws2, rt2, 1, "TOTALE già speso", bold=True, fill=TOT); put(ws2, rt2, 5, f"=SUM(E5:E{rt2 - 1})", bold=True, fmt=EUR, fill=TOT)
    for j in (2, 3, 4, 6): put(ws2, rt2, j, None, fill=TOT)
    ws2.freeze_panes = "A5"
    # ------------------------------------------------------------ Pezzi stampati
    ws3 = wb.create_sheet("Pezzi stampati")
    parts = json.load(open(os.path.join(OUT, "docs", "_pezzi.json")))
    ws3["A1"] = "Pezzi da stampare (file in stl/ e step/, già nell'orientamento di stampa)"; ws3["A1"].font = Font(name=F, bold=True, size=14)
    ws3["A2"] = "Ipotesi per la stima del filamento (celle in blu, modificabili):"; ws3["A2"].font = Font(name=F, size=10, italic=True)
    ws3["A3"] = "Densità PETG (g/cm3)"; ws3["B3"] = 1.27; ws3["C3"] = "Quota di materiale rispetto al pieno (pareti + riempimento 25 %)"; ws3["D3"] = 0.5
    for a in ("A3", "C3"): ws3[a].font = Font(name=F, size=10)
    for a in ("B3", "D3"): ws3[a].font = Font(name=F, size=10, color="0000FF")
    header(ws3, 5, ["Pezzo", "Q.tà", "Volume pieno (cm3)", "Volume totale (cm3)", "Ingombro X (mm)", "Ingombro Y (mm)", "Altezza di stampa Z (mm)", "Note di stampa"], [36, 7, 12, 12, 12, 12, 12, 60])
    for i, p in enumerate(parts):
        r = 6 + i
        put(ws3, r, 1, p["nome"], bold=True); put(ws3, r, 2, p["q"], align="center"); put(ws3, r, 3, p["volume_cm3"], fmt="0.0")
        put(ws3, r, 4, f"=B{r}*C{r}", fmt="0.0")
        for j, v in enumerate(p["ingombro"]): put(ws3, r, 5 + j, v, fmt="0.0")
        put(ws3, r, 8, p["note"])
    rt3 = 6 + len(parts)
    put(ws3, rt3, 1, "TOTALE", bold=True, fill=TOT); put(ws3, rt3, 2, f"=SUM(B6:B{rt3 - 1})", bold=True, fill=TOT, align="center")
    put(ws3, rt3, 4, f"=SUM(D6:D{rt3 - 1})", bold=True, fmt="0", fill=TOT)
    for j in (3, 5, 6, 7, 8): put(ws3, rt3, j, None, fill=TOT)
    put(ws3, rt3 + 1, 1, "Filamento stimato (kg)", bold=True); put(ws3, rt3 + 1, 4, f"=D{rt3}*B3*D3/1000", bold=True, fmt="0.0")
    put(ws3, rt3 + 2, 1, "Ingombro massimo sul piatto (mm)"); put(ws3, rt3 + 2, 5, f"=MAX(E6:E{rt3 - 1})", fmt="0.0"); put(ws3, rt3 + 2, 6, f"=MAX(F6:F{rt3 - 1})", fmt="0.0"); put(ws3, rt3 + 2, 7, f"=MAX(G6:G{rt3 - 1})", fmt="0.0")
    ws3.freeze_panes = "B6"
    # ------------------------------------------------------------ Viteria
    ws4 = wb.create_sheet("Viteria")
    ws4["A1"] = "Viteria: dove va ogni vite"; ws4["A1"].font = Font(name=F, bold=True, size=14)
    ws4["A2"] = "Le viti autofilettanti da 3 mm entrano nei fori Ø2,6 dei pezzi stampati; le lunghezze sono verificate sul modello (verifiche.txt, check_screws.py). Le quantità per tipo sono sommate a destra."; ws4["A2"].font = Font(name=F, size=9, italic=True)
    header(ws4, 4, ["Tipo", "Q.tà", "Dove"], [38, 7, 60])
    for i, (t, q, w) in enumerate(SCREWS):
        r = 5 + i; put(ws4, r, 1, t, bold=True); put(ws4, r, 2, q, align="center"); put(ws4, r, 3, w)
    re_ = 5 + len(SCREWS) - 1
    types = []
    for t, *_ in SCREWS:
        if t not in types: types.append(t)
    ws4.column_dimensions["E"].width = 38; ws4.column_dimensions["F"].width = 10
    for j, c in ((5, "Tipo"), (6, "Totale")):
        cell = ws4.cell(row=4, column=j, value=c); cell.font = Font(name=F, bold=True, color="FFFFFF", size=10); cell.fill = HDR; cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for k, t in enumerate(types):
        put(ws4, 5 + k, 5, t); put(ws4, 5 + k, 6, f"=SUMIF($A$5:$A${re_},E{5 + k},$B$5:$B${re_})", align="center")
    put(ws4, 5 + len(types), 5, "Totale viti", bold=True, fill=TOT); put(ws4, 5 + len(types), 6, f"=SUM(F5:F{4 + len(types)})", bold=True, fill=TOT, align="center")
    ws4.freeze_panes = "A5"
    path = os.path.join(OUT, "docs", "NuovaDistinta.xlsx"); wb.save(path)
    return path

if __name__ == "__main__":
    print(build())
